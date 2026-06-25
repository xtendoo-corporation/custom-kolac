import base64
import importlib.util
import io
import logging
import re
import unicodedata
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from odoo import _, Command, fields, models
from odoo.exceptions import UserError

import openpyxl
import xlrd

_logger = logging.getLogger(__name__)

_ACCOUNT_ASSET_AVAILABLE = importlib.util.find_spec("odoo.addons.account_asset") is not None


class KolacAccountJournalImportWizard(models.TransientModel):
    _name = "kolac.account.journal.import.wizard"
    _description = "Kolac Accounting Import Review"

    DIARY_REQUIRED_COLUMNS = {
        "asiento",
        "fecha",
        "subcuenta",
        "concepto",
        "importe_debe",
        "importe_haber",
    }
    DIARY_OPTIONAL_COLUMNS = {"descripcion", "referencia"}
    DIARY_HEADER_ALIASES = {
        "asiento": "asiento",
        "asto": "asiento",
        "n_asiento": "asiento",
        "num_asiento": "asiento",
        "fecha": "fecha",
        "subcuenta": "subcuenta",
        "cuenta": "subcuenta",
        "cuenta_contable": "subcuenta",
        "concepto": "concepto",
        "importe_debe": "importe_debe",
        "debe": "importe_debe",
        "importe_haber": "importe_haber",
        "haber": "importe_haber",
        "titulo": "descripcion",
        "ttulo": "descripcion",
        "descripcion": "descripcion",
        "descripción": "descripcion",
        "detalle": "descripcion",
        "referencia": "referencia",
        "ref": "referencia",
        "documento": "referencia",
    }
    REFERENCE_HEADER_ALIASES = {
        "subcuenta": "code",
        "cuenta": "code",
        "codigo": "code",
        "código": "code",
        "codigo_cuenta": "code",
        "cuenta_contable": "code",
        "nombre": "name",
        "descripcion": "name",
        "descripción": "name",
        "denominacion": "name",
        "denominación": "name",
        "razon_social": "name",
        "razón_social": "name",
        "titular": "name",
        "nif": "vat",
        "cif": "vat",
        "vat": "vat",
        "notas": "note",
        "observaciones": "note",
        "tipo": "type",
        "grupo": "type",
    }
    PARTNER_PREFIX_MAP = {
        "430": "partner_customer",
        "400": "partner_supplier",
        "410": "partner_creditor",
        "447": "partner_customer",
    }
    MISSING_ACCOUNT_DEFAULTS = {
        "210000": {
            "name": "Terrenos y bienes naturales",
            "account_type": "asset_fixed",
        },
        "211000": {
            "name": "Construcciones",
            "account_type": "asset_fixed",
        },
        "206000": {
            "name": "Aplicaciones informáticas",
            "account_type": "asset_fixed",
        },
        "216000": {
            "name": "Mobiliario",
            "account_type": "asset_fixed",
        },
        "217000": {
            "name": "Equipos para procesos de información",
            "account_type": "asset_fixed",
        },
        "218000": {
            "name": "Elementos de transporte",
            "account_type": "asset_fixed",
        },
        "219000": {
            "name": "Otro inmovilizado material",
            "account_type": "asset_fixed",
        },
        "281100": {
            "name": "Amortización acumulada de construcciones",
            "account_type": "asset_non_current",
        },
        "281200": {
            "name": "Amortización acumulada de instalaciones técnicas",
            "account_type": "asset_non_current",
        },
        "281600": {
            "name": "Amortización acumulada de mobiliario",
            "account_type": "asset_non_current",
        },
        "281700": {
            "name": "Amortización acumulada de equipos para procesos de información",
            "account_type": "asset_non_current",
        },
        "281800": {
            "name": "Amortización acumulada de elementos de transporte",
            "account_type": "asset_non_current",
        },
        "281900": {
            "name": "Amortización acumulada de otro inmovilizado material",
            "account_type": "asset_non_current",
        },
        "412000": {
            "name": "Becarios",
            "account_type": "liability_payable",
            "reconcile": True,
        },
        "447000": {
            "name": "Deudores por subvenciones concedidas",
            "account_type": "asset_receivable",
            "reconcile": True,
        },
        "622060": {
            "name": "Reparación y conservación de equipos",
            "account_type": "expense",
        },
        "629004": {
            "name": "Servicios exteriores - Proveedores Amazon / China",
            "account_type": "expense",
        },
        "629010": {
            "name": "Gastos varios",
            "account_type": "expense",
        },
        "629011": {
            "name": "Trabajos de artes gráficas",
            "account_type": "expense",
        },
        "629013": {
            "name": "Canon impresora",
            "account_type": "expense",
        },
        "629014": {
            "name": "Servicios exteriores - Monitor Informática",
            "account_type": "expense",
        },
        "629040": {
            "name": "Material informático",
            "account_type": "expense",
        },
        "629050": {
            "name": "Telefonía móvil - Movistar",
            "account_type": "expense",
        },
        "629051": {
            "name": "Telefonía móvil - Yoigo",
            "account_type": "expense",
        },
        "629052": {
            "name": "Telefonía móvil - Vodafone",
            "account_type": "expense",
        },
        "629053": {
            "name": "Telefonía móvil - Pepephone",
            "account_type": "expense",
        },
        "629060": {
            "name": "Material de oficina",
            "account_type": "expense",
        },
        "629070": {
            "name": "Material bibliográfico y libros",
            "account_type": "expense",
        },
        "629090": {
            "name": "Gastos de viaje",
            "account_type": "expense",
        },
        "640010": {
            "name": "Salarios becarios",
            "account_type": "expense",
        },
        "650010": {
            "name": "Ayudas monetarias individuales",
            "account_type": "expense",
        },
    }
    EXPLICIT_ACCOUNT_RULES = {
        "206000000": {
            "target_code": "206000",
            "mapping_type": "asset",
        },
        "412000000": {
            "target_code": "412000",
            "mapping_type": "grant_or_becario_payable",
            "warning": "Confirmar con asesor si 412000 Becarios debe mantenerse como cuenta acreedora específica o si debería reclasificarse a 465000/410000. El cliente indica que representa pagos a becarios de proyecto.",
        },
        "622060000": {
            "target_code": "622060",
            "mapping_type": "expense_nature",
        },
        "629004000": {
            "target_code": "629004",
            "mapping_type": "expense_nature_by_supplier",
            "warning": "Proveedor chino/Amazon detectado. Revisar nombre fiscal completo del contacto si se necesita mayor precisión.",
        },
        "629010000": {
            "target_code": "629010",
            "mapping_type": "expense_nature",
        },
        "629010013": {
            "target_code": "629013",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Canon impresora",
        },
        "629010014": {
            "target_code": "629014",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Monitor Informática",
        },
        "629011000": {
            "target_code": "629011",
            "mapping_type": "expense_nature",
        },
        "629040000": {
            "target_code": "629040",
            "mapping_type": "expense_nature",
        },
        "629050000": {
            "target_code": "629050",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Movistar",
        },
        "629050001": {
            "target_code": "629051",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Yoigo",
        },
        "629050002": {
            "target_code": "629052",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Vodafone",
        },
        "629050003": {
            "target_code": "629053",
            "mapping_type": "expense_nature_by_supplier",
            "partner_name": "Pepephone",
        },
        "629060000": {
            "target_code": "629060",
            "mapping_type": "expense_nature",
        },
        "629070000": {
            "target_code": "629070",
            "mapping_type": "expense_nature",
        },
        "629090000": {
            "target_code": "629090",
            "mapping_type": "expense_nature",
        },
        "640010000": {
            "target_code": "640010",
            "mapping_type": "expense_nature",
        },
        "650010000": {
            "target_code": "650010",
            "mapping_type": "expense_nature",
        },
    }
    ASSET_PREFIX_MAP = {
        "211": "211000",
        "206": "206000",
        "216": "216000",
        "217": "217000",
        "218": "218000",
        "219": "219000",
    }
    DEPRECIATION_PREFIX_MAP = {
        "2811": "281100",
        "2812": "281200",
        "2816": "281600",
        "2817": "281700",
        "2818": "281800",
        "2819": "281900",
    }
    ASSET_DEPRECIATION_PREFIX_MAP = {
        "211": "2811",
        "216": "2816",
        "217": "2817",
        "218": "2818",
        "219": "2819",
    }
    TAX_PREFIX_MAP = {
        "472": "tax_input",
        "477": "tax_output",
    }
    OPENING_KEYWORDS = (
        "asiento de apertura",
        "apertura",
        "saldo inicial",
        "saldos iniciales",
        "migracion",
        "migración",
        "cuenta antigua",
        "traspaso",
        "regularizacion inicial",
        "regularización inicial",
    )
    INVOICE_KEYWORDS = (
        "factura",
        "fact",
        "fra",
        "compra",
        "venta",
        "abono",
        "rectificativa",
    )
    MAPPING_TYPES = [
        ("direct_account", "Cuenta directa"),
        ("expense_nature", "Gasto por naturaleza"),
        ("expense_nature_by_supplier", "Gasto por naturaleza y proveedor"),
        ("grant_or_becario_payable", "Acreedor becarios/subvenciones"),
        ("partner_customer", "Cliente"),
        ("partner_supplier", "Proveedor"),
        ("partner_creditor", "Acreedor"),
        ("asset", "Activo"),
        ("accumulated_depreciation", "Amortización acumulada"),
        ("tax_output", "IVA repercutido"),
        ("tax_input", "IVA soportado"),
        ("bank", "Banco"),
        ("cash", "Caja"),
        ("review_required", "Revisión"),
    ]

    diary_file = fields.Binary(string="Fichero Libro Diario", required=True)
    diary_file_name = fields.Char(string="Nombre Libro Diario")
    diary_sheet_name = fields.Char(string="Hoja Libro Diario")
    account_list_file = fields.Binary(string="Fichero Listado de Cuentas")
    account_list_file_name = fields.Char(string="Nombre Listado de Cuentas")
    account_list_sheet_name = fields.Char(string="Hoja Listado")
    account_plan_file = fields.Binary(string="Fichero Plan de Cuentas")
    account_plan_file_name = fields.Char(string="Nombre Plan de Cuentas")
    account_plan_sheet_name = fields.Char(string="Hoja Plan")
    company_id = fields.Many2one(
        "res.company",
        string="Compañía",
        required=True,
        default=lambda self: self.env.company,
    )
    journal_id = fields.Many2one(
        "account.journal",
        string="Diario por defecto",
        required=True,
        domain="[('type', '=', 'general'), ('company_id', '=', company_id)]",
    )
    create_missing_accounts = fields.Boolean(
        string="Permitir crear subcuentas faltantes",
        default=False,
        help="Se podrán crear automáticamente las subcuentas que no existan todavía en Odoo respetando el código original del diario.",
    )
    post_moves = fields.Boolean(string="Publicar asientos válidos", default=False)
    report_year = fields.Integer(
        string="Año del ejercicio",
        help=(
            "Año al que corresponde la importación. Se utiliza para interpretar las "
            "fechas de la columna de fecha del Libro Diario cuando solo indican día y "
            "mes (por ejemplo, '15-Ene.' o '15/01'). Si se deja vacío, se intentará "
            "deducir del propio fichero."
        ),
    )
    state = fields.Selection(
        [("upload", "Carga"), ("review", "Revisión")],
        string="Estado",
        default="upload",
        required=True,
    )
    review_report_file = fields.Binary(string="Informe de revisión")
    review_report_file_name = fields.Char(string="Nombre informe")
    summary_text = fields.Text(string="Resumen", readonly=True)
    detected_move_count = fields.Integer(string="Asientos detectados", readonly=True)
    detected_line_count = fields.Integer(string="Líneas detectadas", readonly=True)
    detected_account_count = fields.Integer(string="Cuentas antiguas detectadas", readonly=True)
    detected_partner_count = fields.Integer(string="Contactos detectados", readonly=True)
    detected_asset_count = fields.Integer(string="Activos detectados", readonly=True)
    detected_tax_count = fields.Integer(string="Impuestos detectados", readonly=True)
    warning_count = fields.Integer(string="Advertencias", readonly=True)
    error_count = fields.Integer(string="Errores", readonly=True)
    mapping_line_ids = fields.One2many(
        "kolac.account.journal.import.map.line",
        "wizard_id",
        string="Mapeos",
    )
    asset_line_ids = fields.One2many(
        "kolac.account.journal.import.asset.line",
        "wizard_id",
        string="Activos",
    )
    move_preview_line_ids = fields.One2many(
        "kolac.account.journal.import.move.preview",
        "wizard_id",
        string="Asientos preparados",
    )
    log_line_ids = fields.One2many(
        "kolac.account.journal.import.log.line",
        "wizard_id",
        string="Logs",
    )
    preview_batch_id = fields.Many2one(
        "kolac.account.import.batch",
        string="Lote de simulación",
        readonly=True,
    )

    def action_analyze(self):
        self.ensure_one()
        if not self.diary_file:
            raise UserError(_("Debe cargar el fichero Libro Diario para iniciar el análisis."))

        self._reset_review_lines()
        analysis = self._analyze_sources()
        batch_hashes = self._build_source_hashes()
        preview_batch = self._upsert_review_batch(analysis, batch_hashes)
        self.write({
            "state": "review",
            "preview_batch_id": preview_batch.id,
            "summary_text": analysis["summary_text"],
            "review_report_file": analysis["review_report_file"],
            "review_report_file_name": analysis["review_report_file_name"],
            "detected_move_count": analysis["detected_move_count"],
            "detected_line_count": analysis["detected_line_count"],
            "detected_account_count": analysis["detected_account_count"],
            "detected_partner_count": analysis["detected_partner_count"],
            "detected_asset_count": analysis["detected_asset_count"],
            "detected_tax_count": analysis["detected_tax_count"],
            "warning_count": analysis["warning_count"],
            "error_count": analysis["error_count"],
            "mapping_line_ids": [Command.create(vals) for vals in analysis["mapping_lines"]],
            "asset_line_ids": [Command.create(vals) for vals in analysis["asset_lines"]],
            "move_preview_line_ids": [Command.create(vals) for vals in analysis["move_preview_lines"]],
            "log_line_ids": [Command.create(vals) for vals in analysis["log_lines"]],
        })
        return self._reopen_wizard()

    def action_confirm_import(self):
        self.ensure_one()
        if self.state != "review":
            return self.action_analyze()
        self._ensure_full_import_requirements()

        blocking_logs = self.log_line_ids.filtered("blocking")
        if blocking_logs:
            raise UserError(
                _("La importación está bloqueada hasta resolver estos problemas:\n%s")
                % "\n".join(blocking_logs.mapped("message")[:20])
            )

        entries, diary_hashes = self._build_import_entries()
        batch = self._create_batch(diary_hashes)
        asset_map = self._ensure_assets(batch)
        created_moves = self._create_moves(entries, batch, asset_map)
        self._save_selected_mappings(batch, asset_map)
        batch.write({
            "state": "done",
            "move_count": len(created_moves),
            "line_count": len(batch.trace_line_ids),
            "imported_at": fields.Datetime.now(),
        })
        if self.post_moves:
            balanced_moves = created_moves.filtered(lambda move: self._is_move_balanced(move))
            if balanced_moves:
                balanced_moves.action_post()
        return self._open_created_moves(created_moves)

    def action_back_to_upload(self):
        self.ensure_one()
        self._reset_review_lines()
        self.write({
            "state": "upload",
            "preview_batch_id": False,
            "summary_text": False,
            "review_report_file": False,
            "review_report_file_name": False,
            "detected_move_count": 0,
            "detected_line_count": 0,
            "detected_account_count": 0,
            "detected_partner_count": 0,
            "detected_asset_count": 0,
            "detected_tax_count": 0,
            "warning_count": 0,
            "error_count": 0,
        })
        return self._reopen_wizard()

    def action_download_review_report(self):
        self.ensure_one()
        if not self.review_report_file:
            raise UserError(_("Primero debe ejecutar el análisis para generar el informe de revisión."))
        return {
            "type": "ir.actions.act_url",
            "url": (
                "/web/content?model=%s&id=%s&field=review_report_file&filename_field=review_report_file_name&download=true"
                % (self._name, self.id)
            ),
            "target": "self",
        }

    def _analyze_sources(self):
        diary_rows = self._load_excel_rows(self.diary_file, self.diary_file_name, self.diary_sheet_name)
        header_row_index, columns = self._locate_diary_columns(diary_rows)
        diary_context = self._get_diary_context(diary_rows, header_row_index, self.diary_file_name)
        diary_lines = self._extract_diary_lines(
            diary_rows,
            header_row_index,
            columns,
            report_year=diary_context["report_year"],
        )

        log_lines = []
        account_list_index = {}
        account_plan_index = {}

        diary_summary = self._summarize_diary_lines(diary_lines)
        mapping_lines = []
        mapping_by_code = {}
        partner_count = 0
        tax_count = 0

        for old_code, summary in diary_summary["accounts"].items():
            mapping_vals = self._build_mapping_line_vals(old_code, summary, account_list_index, account_plan_index, log_lines)
            mapping_lines.append(mapping_vals)
            mapping_by_code[old_code] = mapping_vals
            if mapping_vals.get("partner_name"):
                partner_count += 1
            if mapping_vals.get("target_tax_id"):
                tax_count += 1

        move_preview_lines = self._build_move_preview_lines(diary_summary["entries"], mapping_by_code, log_lines)
        asset_lines = self._build_asset_review_lines(mapping_by_code, diary_summary["accounts"], log_lines)
        error_count = len([log for log in log_lines if log["level"] == "error"])
        warning_count = len([log for log in log_lines if log["level"] == "warning"])
        summary_text = self._build_summary_text(
            diary_summary,
            len(mapping_lines),
            partner_count,
            len(asset_lines),
            tax_count,
            warning_count,
            error_count,
        )
        review_report = self._build_review_report(summary_text, mapping_lines, asset_lines, move_preview_lines, log_lines)
        return {
            "summary_text": summary_text,
            "review_report_file": base64.b64encode(review_report.encode("utf-8")),
            "review_report_file_name": "kolac_import_review.txt",
            "detected_move_count": len(diary_summary["entries"]),
            "detected_line_count": len(diary_lines),
            "detected_account_count": len(mapping_lines),
            "detected_partner_count": partner_count,
            "detected_asset_count": len(asset_lines),
            "detected_tax_count": tax_count,
            "warning_count": warning_count,
            "error_count": error_count,
            "mapping_lines": mapping_lines,
            "asset_lines": asset_lines,
            "move_preview_lines": move_preview_lines,
            "log_lines": log_lines,
        }

    def _build_import_entries(self):
        diary_rows = self._load_excel_rows(self.diary_file, self.diary_file_name, self.diary_sheet_name)
        header_row_index, columns = self._locate_diary_columns(diary_rows)
        diary_context = self._get_diary_context(diary_rows, header_row_index, self.diary_file_name)
        diary_lines = self._extract_diary_lines(
            diary_rows,
            header_row_index,
            columns,
            report_year=diary_context["report_year"],
        )
        entries = OrderedDict()
        mapping_index = {line.source_code: line for line in self.mapping_line_ids}
        preview_index = {line.old_move_number: line for line in self.move_preview_line_ids}

        for line in diary_lines:
            mapping = mapping_index.get(line["old_account_code"])
            if not mapping:
                raise UserError(_("No existe un mapeo revisado para la cuenta antigua %s.") % line["old_account_code"])
            target_account = self._ensure_target_account(mapping)
            partner = self._ensure_partner(mapping)
            preview = preview_index.get(line["old_move_number"])
            if not preview:
                raise UserError(_("No existe una previsualización revisada para el asiento %s.") % line["old_move_number"])

            entry = entries.setdefault(
                line["old_move_number"],
                {
                    "old_move_number": line["old_move_number"],
                    "date": line["date"],
                    "reference": line["reference"],
                    "concept": line["concept"],
                    "journal_id": preview.target_journal_id.id or self.journal_id.id,
                    "lines": [],
                    "fiscal_warnings": [],
                    "is_opening_migration": False,
                },
            )
            entry["lines"].append({
                "name": self._build_line_label(line, mapping),
                "account_id": target_account.id,
                "partner_id": partner.id if partner else False,
                "tax_id": mapping.target_tax_id.id,
                "tax_ids": [],
                "tax_tag_ids": [],
                "tax_repartition_line_id": False,
                "asset_old_code": mapping.source_code if mapping.mapping_type == "asset" else False,
                "old_account_code": line["old_account_code"],
                "old_account_name": mapping.source_account_name,
                "original_label": line["description"] or line["concept"],
                "source_line": line["source_line"],
                "mapping_type": mapping.mapping_type,
                "debit": line["debit"],
                "credit": line["credit"],
                "reference": line["reference"],
            })

        for entry in entries.values():
            entry["is_opening_migration"] = self._is_opening_migration_entry(entry)
            self._annotate_entry_tax_lines(entry)
            self._sanitize_entry_fiscal_metadata(entry)
            total_debit = sum((line["debit"] for line in entry["lines"]), Decimal("0.00"))
            total_credit = sum((line["credit"] for line in entry["lines"]), Decimal("0.00"))
            if self._round_amount(total_debit) != self._round_amount(total_credit):
                raise UserError(
                    _("El asiento %s no está cuadrado y no se importará.") % entry["old_move_number"]
                )

        hashes = self._build_source_hashes()
        return list(entries.values()), hashes

    def _annotate_entry_tax_lines(self, entry):
        if not self._entry_supports_fiscal_taxes(entry):
            return

        tax_lines = [
            line for line in entry["lines"]
            if line["mapping_type"] in {"tax_input", "tax_output"} and line["tax_id"]
        ]
        if not tax_lines:
            return

        for tax_line in tax_lines:
            tax = self.env["account.tax"].browse(tax_line["tax_id"]).exists()
            if not tax:
                continue

            base_repartition_line, tax_repartition_line = self._get_tax_repartition_lines(tax)
            if tax_repartition_line:
                tax_line["tax_repartition_line_id"] = tax_repartition_line.id
                tax_line["tax_tag_ids"] = tax_repartition_line.tag_ids.ids

            for base_line in self._get_entry_tax_base_lines(entry, tax_line):
                base_line["tax_ids"] = sorted(set(base_line["tax_ids"] + [tax.id]))
                if base_repartition_line:
                    base_line["tax_tag_ids"] = sorted(
                        set(base_line["tax_tag_ids"] + base_repartition_line.tag_ids.ids)
                    )

    def _get_entry_tax_base_lines(self, entry, tax_line):
        tax_is_debit = bool(tax_line["debit"])
        return [
            line for line in entry["lines"]
            if line["mapping_type"] not in {"tax_input", "tax_output"}
            and bool(line["debit"]) == tax_is_debit
        ]

    @staticmethod
    def _get_tax_repartition_lines(tax):
        base_repartition_line = tax.invoice_repartition_line_ids.filtered(
            lambda line: line.repartition_type == "base"
        )[:1]
        tax_repartition_line = tax.invoice_repartition_line_ids.filtered(
            lambda line: line.repartition_type == "tax" and line.factor_percent > 0
        )[:1]
        return base_repartition_line, tax_repartition_line

    def _entry_supports_fiscal_taxes(self, entry):
        if entry.get("is_opening_migration"):
            return False

        has_tax_line = any(
            line["mapping_type"] in {"tax_input", "tax_output"} and line["tax_id"]
            for line in entry["lines"]
        )
        has_counterparty = any(
            line["mapping_type"] in {"partner_customer", "partner_supplier", "partner_creditor"}
            for line in entry["lines"]
        )
        return has_tax_line and has_counterparty

    def _is_opening_migration_entry(self, entry):
        if self._is_opening_journal(entry.get("journal_id")):
            return True

        if entry.get("date") == self._get_fiscalyear_start(entry.get("date")):
            return True

        entry_text = self._normalize_keyword_text(
            " ".join(filter(None, [
                entry.get("concept"),
                entry.get("reference"),
                *(line.get("name") for line in entry["lines"]),
                *(line.get("original_label") for line in entry["lines"]),
            ]))
        )
        if any(keyword in entry_text for keyword in self.OPENING_KEYWORDS):
            return True

        if any(line["mapping_type"] in {"asset", "accumulated_depreciation"} for line in entry["lines"]):
            has_invoice_counterparty = any(
                line["mapping_type"] in {"partner_customer", "partner_supplier", "partner_creditor"}
                for line in entry["lines"]
            )
            has_invoice_keyword = any(keyword in entry_text for keyword in self.INVOICE_KEYWORDS)
            if not has_invoice_counterparty or not has_invoice_keyword:
                return True

        return False

    def _is_opening_journal(self, journal_id):
        journal = self.env["account.journal"].browse(journal_id).exists()
        if not journal:
            return False
        journal_text = self._normalize_keyword_text(" ".join(filter(None, [journal.code, journal.name])))
        return any(keyword in journal_text for keyword in ("apertura", "opening", "saldo inicial", "traspaso"))

    def _get_fiscalyear_start(self, entry_date):
        if not entry_date:
            return False
        company = self.company_id
        fiscalyear_last_month = int(company.fiscalyear_last_month or 12)
        fiscalyear_last_day = int(company.fiscalyear_last_day or 31)
        if fiscalyear_last_month == 12 and fiscalyear_last_day == 31:
            return date(entry_date.year, 1, 1)
        if (entry_date.month, entry_date.day) <= (fiscalyear_last_month, fiscalyear_last_day):
            fiscalyear_end = date(entry_date.year, fiscalyear_last_month, fiscalyear_last_day)
        else:
            fiscalyear_end = date(entry_date.year + 1, fiscalyear_last_month, fiscalyear_last_day)
        previous_fiscalyear_end = date(fiscalyear_end.year - 1, fiscalyear_last_month, fiscalyear_last_day)
        return previous_fiscalyear_end + timedelta(days=1)

    def _sanitize_entry_fiscal_metadata(self, entry):
        if self._entry_supports_fiscal_taxes(entry):
            return

        warnings = []
        for line in entry["lines"]:
            had_fiscal_metadata = bool(
                line.get("tax_id")
                or line.get("tax_ids")
                or line.get("tax_tag_ids")
                or line.get("tax_repartition_line_id")
            )
            if not had_fiscal_metadata:
                continue
            line["tax_id"] = False
            line["tax_ids"] = []
            line["tax_tag_ids"] = []
            line["tax_repartition_line_id"] = False
            warnings.append(
                (
                    _(
                        "Se han eliminado impuestos/etiquetas fiscales de una línea de apertura porque los asientos de apertura no deben afectar a modelos fiscales."
                    )
                    if entry.get("is_opening_migration") else
                    _(
                        "Se han eliminado impuestos/etiquetas fiscales de una línea importada porque solo las facturas reales de compra o venta deben afectar a modelos fiscales."
                    )
                )
                + " "
                + _("Asiento %(move)s, cuenta %(account)s, línea %(line)s.") % {
                    "move": entry["old_move_number"],
                    "account": line["old_account_code"],
                    "line": line["source_line"],
                }
            )

        if warnings:
            entry["fiscal_warnings"] = warnings

    @staticmethod
    def _format_fiscal_warnings(entries):
        warnings = []
        for entry in entries:
            warnings.extend(entry.get("fiscal_warnings", []))
        if not warnings:
            return False
        return "\n".join(["", "VALIDACIONES FISCALES"] + [f"- {message}" for message in warnings])

    def _build_source_hashes(self):
        batch_model = self.env["kolac.account.import.batch"]
        return {
            "diary_file_hash": batch_model.build_hash(self._decode_binary_field(self.diary_file)),
            "account_list_file_hash": batch_model.build_hash(self._decode_binary_field(self.account_list_file)),
            "account_plan_file_hash": batch_model.build_hash(self._decode_binary_field(self.account_plan_file)),
        }

    def _upsert_review_batch(self, analysis, hashes):
        batch_model = self.env["kolac.account.import.batch"]
        batch = self._get_reusable_preview_batch()
        values = self._prepare_batch_values(hashes, analysis, dry_run=True)
        if batch:
            batch.write(values)
        else:
            batch = batch_model.create(values)
        batch.write({
            "asset_preview_line_ids": [Command.clear()] + [
                Command.create(self._prepare_batch_asset_line_vals(line_vals))
                for line_vals in analysis["asset_lines"]
            ],
        })
        return batch

    def _get_reusable_preview_batch(self):
        self.ensure_one()
        batch = self.preview_batch_id.exists()
        if batch and batch.dry_run and batch.state == "analyzed" and not batch.trace_line_ids:
            return batch
        return self.env["kolac.account.import.batch"]

    @staticmethod
    def _prepare_batch_asset_line_vals(line_vals):
        vals = {
            "source_code": line_vals.get("source_code"),
            "old_account_name": line_vals.get("old_account_name"),
            "asset_name": line_vals.get("asset_name"),
            "target_account_id": line_vals.get("target_account_id"),
            "depreciation_source_code": line_vals.get("depreciation_source_code"),
            "old_depreciation_account_name": line_vals.get("old_depreciation_account_name"),
            "depreciation_target_account_id": line_vals.get("depreciation_target_account_id"),
            "original_value": line_vals.get("original_value"),
            "accumulated_value": line_vals.get("accumulated_value"),
            "residual_value": line_vals.get("residual_value"),
            "acquisition_date": line_vals.get("acquisition_date"),
            "create_asset": line_vals.get("create_asset"),
            "state": line_vals.get("state"),
            "source_move_number": line_vals.get("source_move_number"),
            "source_reference": line_vals.get("source_reference"),
            "source_label": line_vals.get("source_label"),
            "source_file": line_vals.get("source_file"),
            "warning_message": line_vals.get("warning_message"),
        }
        if _ACCOUNT_ASSET_AVAILABLE:
            vals["asset_id"] = line_vals.get("asset_id")
        return vals

    def _prepare_batch_values(self, hashes, analysis=None, dry_run=False):
        analysis = analysis or {}
        signature = self.env["kolac.account.import.batch"].build_signature(
            hashes["diary_file_hash"],
            hashes["account_list_file_hash"],
            hashes["account_plan_file_hash"],
            self.company_id.id,
        )
        return {
            "name": _("Simulación %s") % fields.Datetime.now() if dry_run else _("Importación %s") % fields.Datetime.now(),
            "company_id": self.company_id.id,
            "journal_id": self.journal_id.id,
            "dry_run": dry_run,
            "create_missing_accounts": self.create_missing_accounts,
            "post_moves": self.post_moves,
            "state": "analyzed",
            "diary_file_name": self.diary_file_name,
            "account_list_file_name": self.account_list_file_name,
            "account_plan_file_name": self.account_plan_file_name,
            "diary_file_hash": hashes["diary_file_hash"],
            "account_list_file_hash": hashes["account_list_file_hash"],
            "account_plan_file_hash": hashes["account_plan_file_hash"],
            "import_signature": signature,
            "mapping_line_count": analysis.get("detected_account_count", len(self.mapping_line_ids)),
            "asset_line_count": analysis.get("detected_asset_count", len(self.asset_line_ids)),
            "warning_count": analysis.get("warning_count", self.warning_count),
            "error_count": analysis.get("error_count", self.error_count),
            "review_report": (
                base64.b64decode(analysis["review_report_file"]).decode("utf-8")
                if analysis.get("review_report_file")
                else (base64.b64decode(self.review_report_file).decode("utf-8") if self.review_report_file else self.summary_text)
            ),
        }

    def _create_batch(self, hashes):
        batch_model = self.env["kolac.account.import.batch"]
        batch = self._get_reusable_preview_batch()
        values = self._prepare_batch_values(hashes, dry_run=False)
        if batch:
            batch.write(values)
            return batch
        batch = batch_model.create(values)
        self.preview_batch_id = batch
        return batch

    def _create_moves(self, entries, batch, asset_map):
        move_model = self.env["account.move"].with_company(self.company_id).with_context(
            default_move_type="entry",
            allowed_company_ids=self.company_id.ids,
            check_move_validity=False,
            skip_invoice_sync=True,
        )
        move_line_model = self.env["account.move.line"].with_company(self.company_id).with_context(
            allowed_company_ids=self.company_id.ids,
            check_move_validity=False,
            skip_invoice_sync=True,
        )
        created_moves = self.env["account.move"]

        for entry in entries:
            move = move_model.create({
                "move_type": "entry",
                "company_id": self.company_id.id,
                "journal_id": entry["journal_id"],
                "date": entry["date"],
                "ref": entry["reference"] or entry["concept"] or _("Asiento origen %s") % entry["old_move_number"],
                "narration": _("Importado desde KOLAC. Asiento origen: %s") % entry["old_move_number"],
            })
            created_moves |= move
            created_move_lines = self.env["account.move.line"]
            for source_line in entry["lines"]:
                line_vals = {
                    "move_id": move.id,
                    "name": source_line["name"],
                    "account_id": source_line["account_id"],
                    "partner_id": source_line["partner_id"],
                    "debit": float(source_line["debit"]),
                    "credit": float(source_line["credit"]),
                }
                if source_line["tax_ids"]:
                    line_vals["tax_ids"] = [Command.set(source_line["tax_ids"])]
                if source_line["tax_tag_ids"]:
                    line_vals["tax_tag_ids"] = [Command.set(source_line["tax_tag_ids"])]
                if source_line["tax_repartition_line_id"]:
                    line_vals["tax_repartition_line_id"] = source_line["tax_repartition_line_id"]
                created_move_lines |= move_line_model.create(line_vals)
            for move_line, source_line in zip(created_move_lines, entry["lines"]):
                trace_vals = {
                    "batch_id": batch.id,
                    "move_id": move.id,
                    "move_line_id": move_line.id,
                    "old_move_number": entry["old_move_number"],
                    "old_move_ref": entry["reference"],
                    "move_date": entry["date"],
                    "old_account_code": source_line["old_account_code"],
                    "old_account_name": source_line["old_account_name"],
                    "original_label": source_line["original_label"],
                    "target_account_id": move_line.account_id.id,
                    "partner_id": move_line.partner_id.id,
                    "tax_id": source_line["tax_id"],
                    "mapping_type": source_line["mapping_type"],
                    "source_file": self.diary_file_name,
                    "source_line": source_line["source_line"],
                    "debit": float(source_line["debit"]),
                    "credit": float(source_line["credit"]),
                    "note": source_line["reference"],
                }
                if _ACCOUNT_ASSET_AVAILABLE and source_line["asset_old_code"]:
                    trace_vals["asset_id"] = asset_map.get(source_line["asset_old_code"], self.env["account.asset"]).id or False
                self.env["kolac.account.import.trace"].create(trace_vals)

        fiscal_warning_report = self._format_fiscal_warnings(entries)
        if fiscal_warning_report:
            batch.write({
                "review_report": (batch.review_report or "") + fiscal_warning_report,
                "warning_count": (batch.warning_count or 0) + len([message for entry in entries for message in entry.get("fiscal_warnings", [])]),
            })

        return created_moves

    def _ensure_assets(self, batch):
        asset_map = {}
        if "account.asset" not in self.env:
            return asset_map

        for line in self.asset_line_ids.filtered(lambda asset: asset.create_asset):
            existing_asset = line.asset_id or self._find_saved_mapping(line.source_code).asset_id
            if existing_asset:
                self._write_asset_traceability(existing_asset, line, batch)
                asset_map[line.source_code] = existing_asset
                continue
            if (
                not line.target_account_id
                or not line.depreciation_target_account_id
                or line.original_value <= 0
                or not line.acquisition_date
            ):
                continue
            model = self._find_asset_model_for_line(line)
            if not model:
                continue
            asset = self.env["account.asset"].create({
                "model_id": model.id,
                "name": line.asset_name,
                "company_id": self.company_id.id,
                "account_asset_id": line.target_account_id.id,
                "account_depreciation_id": line.depreciation_target_account_id.id,
                "account_depreciation_expense_id": model.account_depreciation_expense_id.id,
                "journal_id": model.journal_id.id,
                "original_value": float(line.original_value),
                "already_depreciated_amount_import": float(line.accumulated_value),
                "acquisition_date": line.acquisition_date or fields.Date.today(),
                "prorata_date": line.acquisition_date or fields.Date.today(),
            })
            self._write_asset_traceability(asset, line, batch)
            line.asset_id = asset.id
            asset_map[line.source_code] = asset

        return asset_map

    def _find_asset_model_for_line(self, line):
        asset_model = self.env["account.asset"]
        direct_model = asset_model.search([
            ("state", "=", "model"),
            ("company_id", "=", self.company_id.id),
            ("account_asset_id", "=", line.target_account_id.id),
            ("account_depreciation_id", "=", line.depreciation_target_account_id.id),
        ], limit=1)
        if direct_model:
            return direct_model

        template_asset_account = self._find_model_template_account(line.target_account_id.code)
        template_depreciation_account = self._find_model_template_account(line.depreciation_target_account_id.code)
        if not template_asset_account or not template_depreciation_account:
            return asset_model

        return asset_model.search([
            ("state", "=", "model"),
            ("company_id", "=", self.company_id.id),
            ("account_asset_id", "=", template_asset_account.id),
            ("account_depreciation_id", "=", template_depreciation_account.id),
        ], limit=1)

    def _find_model_template_account(self, code):
        for candidate_code in self._get_template_code_candidates(code):
            candidate = self._find_account(candidate_code)
            if candidate:
                return candidate
        return self.env["account.account"]

    def _write_asset_traceability(self, asset, line, batch):
        asset.write({
            "kolac_old_asset_account_code": line.source_code,
            "kolac_old_asset_account_name": line.old_account_name,
            "kolac_old_depreciation_account_code": line.depreciation_source_code,
            "kolac_old_depreciation_account_name": line.old_depreciation_account_name,
            "kolac_original_value_detected": float(line.original_value or 0.0),
            "kolac_accumulated_depreciation_detected": float(line.accumulated_value or 0.0),
            "kolac_net_book_value_detected": float(line.residual_value or 0.0),
            "kolac_acquisition_date_detected": line.acquisition_date,
            "kolac_import_batch_id": batch.id,
            "kolac_source_file": line.source_file,
            "kolac_source_move_number": line.source_move_number,
            "kolac_source_move_ref": line.source_reference,
            "kolac_source_label": line.source_label,
            "kolac_asset_review_state": line.state,
            "kolac_warning_message": line.warning_message,
        })

    def _save_selected_mappings(self, batch, asset_map):
        map_model = self.env["kolac.account.subaccount.map"]
        for line in self.mapping_line_ids.filtered(lambda record: record.save_mapping and record.target_account_id):
            mapping = line.saved_mapping_id or self._find_saved_mapping(line.source_code)
            values = {
                "company_id": self.company_id.id,
                "source_code": line.source_code,
                "source_name": line.source_account_name,
                "source_origin": line.source_origin,
                "target_account_id": line.target_account_id.id,
                "mapping_type": line.mapping_type,
                "partner_id": line.partner_id.id,
                "tax_id": line.target_tax_id.id,
                "action": line.action,
                "state": "manual" if line.review_required else "auto",
                "review_required": line.review_required,
                "warning_message": line.warning_message,
                "import_batch_id": batch.id,
            }
            if _ACCOUNT_ASSET_AVAILABLE:
                resolved_asset = asset_map.get(line.source_code, getattr(line, "asset_id", False))
                values["asset_id"] = resolved_asset.id if resolved_asset else False
            if mapping:
                mapping.write(values)
            else:
                map_model.create(values)

    def _ensure_full_import_requirements(self):
        missing = []
        if not self.diary_file:
            missing.append(_("Libro Diario"))
        if missing:
            raise UserError(
                _("La importación definitiva requiere al menos el Libro Diario. Faltan: %s") % ", ".join(missing)
            )

    def _ensure_target_account(self, mapping_line):
        if mapping_line.target_account_id:
            return mapping_line.target_account_id
        if not self._can_auto_create_missing_account(mapping_line.target_account_code):
            raise UserError(
                _("La cuenta final %s no existe y no está autorizado crearla automáticamente.")
                % mapping_line.target_account_code
            )
        create_vals = self._get_missing_account_vals(
            mapping_line.target_account_code,
            mapping_line.source_account_name,
            mapping_type=mapping_line.mapping_type,
        )
        if not create_vals:
            raise UserError(
                _("No se ha podido inferir cómo crear la cuenta final %s.") % mapping_line.target_account_code
            )
        account = self._get_account_model().create(create_vals)
        mapping_line.target_account_id = account.id
        return account

    def _ensure_partner(self, mapping_line):
        if not self._mapping_supports_partner(mapping_line.mapping_type):
            return self.env["res.partner"]
        if mapping_line.partner_id:
            return mapping_line.partner_id
        if not mapping_line.partner_name:
            if not self._mapping_requires_partner(mapping_line.mapping_type):
                return self.env["res.partner"]
            raise UserError(
                _("La cuenta antigua %s requiere contacto y no tiene nombre identificable.") % mapping_line.source_code
            )
        partner_candidates = self.env["res.partner"].with_company(self.company_id).search([
            ("name", "=", mapping_line.partner_name),
            ("company_id", "in", [False, self.company_id.id]),
        ], limit=2)
        if len(partner_candidates) > 1:
            raise UserError(
                _("El contacto %s es ambiguo. Revise el mapeo antes de importar.") % mapping_line.partner_name
            )
        if partner_candidates:
            mapping_line.partner_id = partner_candidates.id
            return partner_candidates
        partner = self.env["res.partner"].sudo().with_company(self.company_id).create({
            "name": mapping_line.partner_name,
            "company_id": self.company_id.id,
            "company_type": "company",
            "customer_rank": 1 if mapping_line.mapping_type == "partner_customer" else 0,
            "supplier_rank": 1 if mapping_line.mapping_type in {"partner_supplier", "partner_creditor", "expense_nature_by_supplier", "grant_or_becario_payable"} else 0,
        })
        mapping_line.partner_id = partner.id
        return partner

    def _build_line_label(self, source_line, mapping_line):
        label = source_line["description"] or source_line["concept"] or mapping_line.source_account_name or mapping_line.source_code
        if mapping_line.mapping_type in {"asset", "accumulated_depreciation"}:
            return _("%s - cuenta antigua %s") % (label, mapping_line.source_code)
        return label

    def _load_reference_index(self, file_content, file_name, sheet_name, source_origin, log_lines):
        if not file_content:
            return {}
        rows = self._load_excel_rows(file_content, file_name, sheet_name)
        if not rows:
            log_lines.append(self._make_log_line("warning", source_origin, _("El fichero %s está vacío.") % file_name))
            return {}
        columns = self._locate_reference_columns(rows)
        if not columns.get("code"):
            log_lines.append(
                self._make_log_line(
                    "warning",
                    source_origin,
                    _("No se ha encontrado una columna de código reconocible en %s.") % file_name,
                )
            )
            return {}
        index = {}
        for row_number, row in enumerate(rows[1:], start=2):
            old_code = self._sanitize_old_account_code(self._get_value(row, columns, "code"))
            if not old_code:
                continue
            index[old_code] = {
                "name": self._cell_to_text(self._get_value(row, columns, "name")),
                "vat": self._cell_to_text(self._get_value(row, columns, "vat")),
                "note": self._cell_to_text(self._get_value(row, columns, "note")),
                "type": self._cell_to_text(self._get_value(row, columns, "type")),
                "source_origin": source_origin,
                "row_number": row_number,
            }
        return index

    def _locate_reference_columns(self, rows):
        columns = {}
        for row in rows[:5]:
            columns = {}
            for cell_index, value in enumerate(row):
                key = self.REFERENCE_HEADER_ALIASES.get(self._normalize_header(value))
                if key and key not in columns:
                    columns[key] = cell_index
            if columns.get("code"):
                return columns
        return columns

    def _extract_diary_lines(self, rows, header_row_index, columns, report_year=False):
        lines = []
        data_started = False
        for row_number, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
            if self._is_blank_row(row):
                if data_started:
                    break
                continue
            data_started = True
            move_number = self._cell_to_text(self._get_value(row, columns, "asiento"))
            raw_date_value = self._get_value(row, columns, "fecha")
            entry_date = self._parse_date(
                raw_date_value,
                default_year=report_year,
            )
            old_account_code = self._sanitize_old_account_code(self._get_value(row, columns, "subcuenta"))
            concept = self._cell_to_text(self._get_value(row, columns, "concepto"))
            description = self._cell_to_text(self._get_value(row, columns, "descripcion"))
            reference = self._cell_to_text(self._get_value(row, columns, "referencia"))
            debit = self._parse_amount(self._get_value(row, columns, "importe_debe"))
            credit = self._parse_amount(self._get_value(row, columns, "importe_haber"))

            if self._is_zero_amount_row(debit, credit):
                continue
            if self._is_summary_row(row, entry_date, old_account_code):
                continue
            if not any([move_number, entry_date, old_account_code, concept, description, debit, credit]):
                continue
            if not move_number or not entry_date or not old_account_code:
                if not entry_date and raw_date_value not in (None, False, ""):
                    raise UserError(_(
                        "No se ha podido interpretar la fecha '%(value)s' de la fila "
                        "%(row)s del Libro Diario. Indique el 'Año del ejercicio' en el "
                        "asistente de importación para resolver fechas sin año."
                    ) % {
                        "value": self._cell_to_text(raw_date_value),
                        "row": row_number,
                    })
                raise UserError(_("La fila %s del Libro Diario carece de datos obligatorios.") % row_number)
            if debit and credit:
                raise UserError(_("La fila %s informa Debe y Haber simultáneamente.") % row_number)
            if not debit and not credit:
                raise UserError(_("La fila %s no informa Debe ni Haber.") % row_number)

            lines.append({
                "source_line": row_number,
                "old_move_number": move_number,
                "date": entry_date,
                "old_account_code": old_account_code,
                "concept": concept,
                "description": description,
                "reference": reference,
                "debit": debit,
                "credit": credit,
            })
        if not lines:
            raise UserError(_("No se han encontrado líneas válidas en el Libro Diario."))
        return lines

    def _requires_reference_files(self):
        return False

    def _get_diary_context(self, rows, header_row_index, file_name=False):
        diary_format = self._detect_diary_format(rows[header_row_index])
        report_year = self._get_effective_report_year(rows, file_name=file_name)
        return {
            "format": diary_format,
            "report_year": report_year,
            "allows_single_file_import": diary_format == "kolac_official",
        }

    def _get_effective_report_year(self, rows, file_name=False):
        if self.report_year:
            if not 1900 <= self.report_year <= 2200:
                raise UserError(
                    _("El 'Año del ejercicio' indicado (%s) no es válido.") % self.report_year
                )
            return self.report_year
        return self._guess_report_year(rows, file_name=file_name)

    def _detect_diary_format(self, header_row):
        normalized_headers = {
            self._normalize_header(value)
            for value in header_row
            if value not in (None, False, "")
        }
        official_headers = {"fecha", "asto", "ord", "dia", "cuenta", "concepto", "debe", "haber"}
        if official_headers.issubset(normalized_headers) and ({"ttulo"} <= normalized_headers or {"titulo"} <= normalized_headers):
            return "kolac_official"
        return "standard"

    def _guess_report_year(self, rows, file_name=False):
        candidates = []
        for row in rows[:10]:
            text = " ".join(
                self._cell_to_text(value)
                for value in row
                if value not in (None, False, "")
            )
            candidates.extend(re.findall(r"\b(20\d{2})\b", text))
        if not candidates and file_name:
            candidates = re.findall(r"\b(20\d{2})\b", file_name)
        return int(candidates[-1]) if candidates else False

    @staticmethod
    def _decode_binary_field(binary_value):
        if not binary_value:
            return b""
        return base64.b64decode(binary_value)

    def _summarize_diary_lines(self, diary_lines):
        accounts = OrderedDict()
        entries = OrderedDict()
        for line in diary_lines:
            account_summary = accounts.setdefault(line["old_account_code"], {
                "count": 0,
                "debit": Decimal("0.00"),
                "credit": Decimal("0.00"),
                "samples": [],
            })
            account_summary["count"] += 1
            account_summary["debit"] += line["debit"]
            account_summary["credit"] += line["credit"]
            if len(account_summary["samples"]) < 5:
                account_summary["samples"].append(line)

            entry_summary = entries.setdefault(line["old_move_number"], {
                "old_move_number": line["old_move_number"],
                "date": line["date"],
                "reference": line["reference"],
                "concept": line["concept"],
                "count": 0,
                "debit": Decimal("0.00"),
                "credit": Decimal("0.00"),
                "account_codes": set(),
                "lines": [],
            })
            entry_summary["count"] += 1
            entry_summary["debit"] += line["debit"]
            entry_summary["credit"] += line["credit"]
            entry_summary["account_codes"].add(line["old_account_code"])
            entry_summary["lines"].append(line)
        return {"accounts": accounts, "entries": entries}

    def _build_mapping_line_vals(self, old_code, summary, account_list_index, account_plan_index, log_lines):
        saved_mapping = self._find_saved_mapping(old_code)
        account_list_data = account_list_index.get(old_code, {})
        account_plan_data = account_plan_index.get(old_code, {})
        saved_target_account = saved_mapping.target_account_id.filtered(lambda account: account.code == old_code)
        saved_target_code = saved_target_account.code if saved_target_account else False
        source_name = (
            saved_mapping.source_name
            or account_list_data.get("name")
            or account_plan_data.get("name")
            or self._guess_name_from_samples(summary["samples"])
        )
        explicit_rule = self._get_explicit_account_rule(old_code, summary, source_name)
        mapping_type = saved_mapping.mapping_type or explicit_rule.get("mapping_type") or self._detect_mapping_type(old_code)
        target_account_code = saved_target_code or self._map_old_code_to_target_code(old_code)
        target_account = saved_target_account or self._find_account(target_account_code)
        if not target_account and self.create_missing_accounts and self._can_auto_create_missing_account(target_account_code):
            target_account = self._create_missing_account(
                target_account_code,
                source_name,
                mapping_type=mapping_type,
            )
        partner_name = (
            saved_mapping.partner_id.name
            or explicit_rule.get("partner_name")
            or self._guess_partner_name(mapping_type, source_name, summary["samples"])
        )
        partner_id, partner_warning = self._find_partner_candidate(saved_mapping.partner_id, partner_name)
        target_tax = saved_mapping.tax_id or self._find_tax_candidate(mapping_type, source_name, summary["samples"])
        warning_parts = []
        review_required = False
        action = saved_mapping.action or (
            "create_partner"
            if self._mapping_supports_partner(mapping_type) and partner_name and not partner_id
            else "reuse"
        )
        source_origin = saved_mapping.source_origin or account_list_data.get("source_origin") or account_plan_data.get("source_origin") or "journal"

        if explicit_rule.get("warning"):
            warning_parts.append(_(explicit_rule["warning"]))

        if not target_account:
            if self._can_auto_create_missing_account(target_account_code):
                warning_parts.append(_("La cuenta final %s no existe en Odoo y se creará durante la importación.") % target_account_code)
            else:
                warning_parts.append(_("No existe la cuenta final %s en Odoo.") % target_account_code)
                review_required = True
        if self._mapping_requires_partner(mapping_type) and not partner_name:
            warning_parts.append(_("No se ha podido identificar el contacto para la subcuenta %s.") % old_code)
            review_required = True
        if self._mapping_requires_partner(mapping_type) and partner_warning:
            warning_parts.append(partner_warning)
            review_required = True
        if mapping_type == "expense_nature_by_supplier" and partner_warning:
            warning_parts.append(partner_warning)
        if old_code == "629004000" and not partner_name:
            warning_parts.append(_("Proveedor chino/Amazon detectado. No se ha podido identificar con precisión el contacto para la subcuenta %s.") % old_code)
        if mapping_type in {"tax_input", "tax_output"} and not target_tax:
            warning_parts.append(_("No se ha podido asociar un impuesto Odoo fiable para %s.") % old_code)
        if mapping_type == "review_required":
            warning_parts.append(_("La subcuenta %s requiere revisión manual.") % old_code)
            review_required = True

        for warning in warning_parts:
            log_lines.append(self._make_log_line(
                "warning" if mapping_type != "review_required" else "error",
                "mapping",
                warning,
                old_account_code=old_code,
                blocking=review_required and mapping_type == "review_required",
            ))

        vals = {
            "source_code": old_code,
            "source_account_name": source_name,
            "source_origin": source_origin,
            "occurrence_count": summary["count"],
            "saved_mapping_id": saved_mapping.id,
            "mapping_type": mapping_type,
            "target_account_code": target_account_code,
            "target_account_id": target_account.id,
            "partner_name": partner_name,
            "partner_id": partner_id.id,
            "target_tax_id": target_tax.id,
            "suggestion_origin": self._guess_suggestion_origin(saved_mapping, target_account, old_code),
            "save_mapping": True,
            "action": action,
            "review_required": review_required,
            "warning_message": "\n".join(warning_parts),
            "plan_account_name": account_plan_data.get("name"),
            "account_list_name": account_list_data.get("name"),
            "debit_total": float(summary["debit"]),
            "credit_total": float(summary["credit"]),
        }
        if _ACCOUNT_ASSET_AVAILABLE:
            vals["asset_id"] = getattr(saved_mapping, "asset_id", False) and saved_mapping.asset_id.id or False
        return vals

    def _build_move_preview_lines(self, entries, mapping_by_code, log_lines):
        preview_lines = []
        for entry in entries.values():
            balanced = self._round_amount(entry["debit"]) == self._round_amount(entry["credit"])
            target_journal = self._suggest_journal_for_entry(entry, mapping_by_code)
            if not balanced:
                log_lines.append(self._make_log_line(
                    "error",
                    "entry",
                    _("El asiento %s está descuadrado y no podrá importarse.") % entry["old_move_number"],
                    move_number=entry["old_move_number"],
                    blocking=True,
                ))
            preview_lines.append({
                "old_move_number": entry["old_move_number"],
                "move_date": entry["date"],
                "reference": entry["reference"] or entry["concept"],
                "concept": entry["concept"],
                "line_count": entry["count"],
                "debit_total": float(entry["debit"]),
                "credit_total": float(entry["credit"]),
                "target_journal_id": target_journal.id or self.journal_id.id,
                "is_balanced": balanced,
                "status": "ready" if balanced else "error",
            })
        return preview_lines

    def _build_asset_review_lines(self, mapping_by_code, account_summaries, log_lines):
        asset_lines = []
        depreciation_candidates = []
        for old_code, mapping in mapping_by_code.items():
            if mapping["mapping_type"] == "accumulated_depreciation":
                depreciation_candidates.append((old_code, mapping))

        matched_depreciation_codes = set()

        for old_code, mapping in mapping_by_code.items():
            if mapping["mapping_type"] != "asset":
                continue
            depreciation = self._find_matching_depreciation(
                old_code,
                mapping,
                account_summaries,
                depreciation_candidates,
            )
            depreciation_code = depreciation[0] if depreciation else False
            depreciation_mapping = depreciation[1] if depreciation else {}
            if depreciation_code:
                matched_depreciation_codes.add(depreciation_code)
            original_value = max(
                account_summaries[old_code]["debit"] - account_summaries[old_code]["credit"],
                Decimal("0.00"),
            )
            accumulated_value = Decimal("0.00")
            warnings = []
            source_sample = self._pick_primary_sample(account_summaries[old_code]["samples"])
            expected_depreciation_target_code = self._get_expected_depreciation_target_code(old_code)
            depreciation_target_account_id = depreciation_mapping.get("target_account_id")
            if not depreciation_target_account_id and expected_depreciation_target_code:
                depreciation_target = self._find_account(expected_depreciation_target_code)
                if not depreciation_target and self.create_missing_accounts and self._can_auto_create_missing_account(expected_depreciation_target_code):
                    depreciation_target = self._create_missing_account(
                        expected_depreciation_target_code,
                        depreciation_mapping.get("source_account_name") or mapping.get("source_account_name") or old_code,
                        mapping_type="accumulated_depreciation",
                    )
                depreciation_target_account_id = depreciation_target.id
            state = "mapped_with_asset"
            if depreciation_code:
                accumulated_value = max(
                    account_summaries[depreciation_code]["credit"] - account_summaries[depreciation_code]["debit"],
                    Decimal("0.00"),
                )
            else:
                warnings.append(
                    _(
                        "No se ha encontrado amortización acumulada asociada para la cuenta antigua %(code)s. Confirmar si el importe detectado corresponde al valor original del activo o al valor neto pendiente de amortizar."
                    ) % {"code": old_code}
                )
                log_lines.append(self._make_log_line(
                    "warning",
                    "asset",
                    _("El activo %s no tiene amortización acumulada asociada.") % old_code,
                    old_account_code=old_code,
                ))
                state = "review_required"
            residual_value = original_value - accumulated_value
            if residual_value < 0:
                warnings.append(
                    _(
                        "La amortización acumulada detectada supera el valor original del activo. Revisar antes de importar automáticamente."
                    )
                )
                log_lines.append(self._make_log_line(
                    "warning",
                    "asset",
                    _("La amortización acumulada supera el valor del activo %s.") % old_code,
                    old_account_code=old_code,
                ))
                state = "error"
            elif residual_value == 0 and original_value:
                warnings.append(_("Activo totalmente amortizado."))
                state = "fully_depreciated"

            if not mapping["target_account_id"]:
                warnings.append(_("No se ha podido localizar la cuenta Odoo del activo %s.") % old_code)
                state = "review_required"
            if expected_depreciation_target_code and not depreciation_target_account_id:
                warnings.append(
                    _(
                        "No existe la cuenta Odoo de amortización %(code)s para el activo %(asset)s."
                    ) % {"code": expected_depreciation_target_code, "asset": old_code}
                )
                state = "review_required"
            if not source_sample.get("date"):
                warnings.append(_("No se ha podido detectar la fecha del asiento de apertura del activo."))
                state = "review_required"

            create_asset = bool(
                state in {"mapped_with_asset", "fully_depreciated"}
                and mapping["target_account_id"]
                and depreciation_target_account_id
                and original_value > 0
                and source_sample.get("date")
            )

            asset_lines.append({
                "source_code": old_code,
                "asset_name": mapping["source_account_name"],
                "old_account_name": mapping["source_account_name"],
                "target_account_id": mapping["target_account_id"],
                "depreciation_source_code": depreciation_code,
                "old_depreciation_account_name": depreciation_mapping.get("source_account_name"),
                "depreciation_target_account_id": depreciation_target_account_id,
                "original_value": float(original_value),
                "accumulated_value": float(accumulated_value),
                "residual_value": float(residual_value),
                "acquisition_date": source_sample.get("date"),
                "create_asset": create_asset,
                "state": state,
                "warning_message": "\n".join(warnings),
                "source_move_number": source_sample.get("old_move_number"),
                "source_reference": source_sample.get("reference"),
                "source_label": source_sample.get("description") or source_sample.get("concept"),
                "source_file": self.diary_file_name,
            })
            if _ACCOUNT_ASSET_AVAILABLE:
                asset_lines[-1]["asset_id"] = mapping.get("asset_id") or False

        for old_code, mapping in mapping_by_code.items():
            if mapping["mapping_type"] != "accumulated_depreciation":
                continue
            if old_code not in matched_depreciation_codes:
                log_lines.append(self._make_log_line(
                    "warning",
                    "asset",
                    _("Existe amortización acumulada %s sin activo relacionado.") % old_code,
                    old_account_code=old_code,
                ))

        return asset_lines

    def _guess_name_from_samples(self, samples):
        for sample in samples:
            if sample["description"]:
                return sample["description"]
        for sample in samples:
            if sample["concept"]:
                return sample["concept"]
        return False

    def _guess_partner_name(self, mapping_type, source_name, samples):
        if not self._mapping_supports_partner(mapping_type):
            return False
        if source_name:
            return source_name
        return self._guess_name_from_samples(samples)

    @staticmethod
    def _mapping_requires_partner(mapping_type):
        return mapping_type in {"partner_customer", "partner_supplier", "partner_creditor"}

    @classmethod
    def _mapping_supports_partner(cls, mapping_type):
        return cls._mapping_requires_partner(mapping_type) or mapping_type in {
            "expense_nature_by_supplier",
            "grant_or_becario_payable",
        }

    def _find_partner_candidate(self, saved_partner, partner_name):
        if saved_partner:
            return saved_partner, False
        if not partner_name:
            return self.env["res.partner"], False
        partners = self.env["res.partner"].with_company(self.company_id).search([
            ("name", "=", partner_name),
            ("company_id", "in", [False, self.company_id.id]),
        ], limit=2)
        if len(partners) > 1:
            return self.env["res.partner"], _("El contacto %s es ambiguo.") % partner_name
        return partners[:1], False

    def _find_tax_candidate(self, mapping_type, source_name, samples):
        if mapping_type not in {"tax_input", "tax_output"}:
            return self.env["account.tax"]
        text_parts = [source_name or ""] + [sample["description"] or sample["concept"] or "" for sample in samples]
        rate = self._extract_tax_rate(" ".join(text_parts))
        if rate is None:
            return self.env["account.tax"]
        tax_type = "purchase" if mapping_type == "tax_input" else "sale"
        taxes = self.env["account.tax"].with_company(self.company_id).search([
            ("type_tax_use", "in", [tax_type, "none"]),
            ("amount_type", "=", "percent"),
            ("amount", "=", rate),
            ("company_id", "=", self.company_id.id),
        ], limit=1)
        return taxes

    def _extract_tax_rate(self, text):
        if not text:
            return None
        match = re.search(r"(21|10|4)\s*%", text)
        if match:
            return float(match.group(1))
        return None

    def _suggest_journal_for_entry(self, entry, mapping_by_code):
        codes = entry["account_codes"]
        if any(mapping_by_code[code]["mapping_type"] == "cash" for code in codes if code in mapping_by_code):
            journal = self.env["account.journal"].search([
                ("company_id", "=", self.company_id.id),
                ("type", "=", "cash"),
            ], limit=1)
            if journal:
                return journal
        if any(mapping_by_code[code]["mapping_type"] == "bank" for code in codes if code in mapping_by_code):
            journal = self.env["account.journal"].search([
                ("company_id", "=", self.company_id.id),
                ("type", "=", "bank"),
            ], limit=1)
            if journal:
                return journal
        return self.journal_id

    def _guess_suggestion_origin(self, saved_mapping, target_account, old_code):
        if saved_mapping:
            return "saved"
        if target_account and target_account.code == old_code:
            return "suggested"
        if old_code in self.EXPLICIT_ACCOUNT_RULES:
            return "suggested"
        return "none"

    def _detect_mapping_type(self, old_code):
        if any(old_code.startswith(prefix) for prefix in self.PARTNER_PREFIX_MAP):
            return self.PARTNER_PREFIX_MAP[old_code[:3]]
        if old_code[:3] in self.ASSET_PREFIX_MAP:
            return "asset"
        if old_code[:4] in self.DEPRECIATION_PREFIX_MAP:
            return "accumulated_depreciation"
        if any(old_code.startswith(prefix) for prefix in self.TAX_PREFIX_MAP):
            return self.TAX_PREFIX_MAP[old_code[:3]]
        if old_code.startswith("572"):
            return "bank"
        if old_code.startswith("570"):
            return "cash"
        if not old_code.isdigit():
            return "review_required"
        return "direct_account"

    def _map_old_code_to_target_code(self, old_code):
        return old_code

    def _map_special_group_code(self, old_code):
        return False

    def _get_explicit_account_rule(self, old_code, summary, source_name):
        del summary, source_name
        return dict(self.EXPLICIT_ACCOUNT_RULES.get(old_code, {}))

    def _get_material_it_asset_rule(self, samples, source_name):
        text = self._normalize_keyword_text(
            " ".join(filter(None, [source_name] + [sample["description"] or sample["concept"] or "" for sample in samples]))
        )
        amount = self._get_largest_line_amount(samples)
        warning = _(
            "Revisar si corresponde a gasto 629040 o inmovilizado 206000/216000/217000 según criterio de 300 euros de base imponible."
        )
        family_map = {
            "206000": ("software", "aplicacion", "aplicación", "programa", "licencia"),
            "216000": ("mobiliario", "mesa", "silla", "armario", "estanteria", "estantería"),
            "217000": ("equipo informatico", "equipo informático", "ordenador", "portatil", "portátil", "impresora", "monitor", "servidor", "informat"),
        }
        for target_code, keywords in family_map.items():
            if any(keyword in text for keyword in keywords):
                if amount >= Decimal("300.00"):
                    return {
                        "target_code": target_code,
                        "mapping_type": "asset",
                    }
                return {}
        if amount >= Decimal("300.00"):
            return {
                "target_code": "629040",
                "mapping_type": "expense_nature",
                "warning": warning,
            }
        return {}

    @staticmethod
    def _get_largest_line_amount(samples):
        amounts = [max(sample["debit"], sample["credit"]) for sample in samples]
        return max(amounts) if amounts else Decimal("0.00")

    def _find_matching_depreciation(self, asset_code, asset_mapping, account_summaries, depreciation_candidates):
        asset_summary = account_summaries[asset_code]
        expected_prefix = self.ASSET_DEPRECIATION_PREFIX_MAP.get(asset_code[:3])
        best_match = False
        best_score = -1
        for depreciation_code, depreciation_mapping in depreciation_candidates:
            if expected_prefix and not depreciation_code.startswith(expected_prefix):
                continue
            score = self._score_depreciation_match(
                asset_code,
                asset_mapping,
                asset_summary,
                depreciation_code,
                depreciation_mapping,
                account_summaries[depreciation_code],
            )
            if score > best_score:
                best_score = score
                best_match = (depreciation_code, depreciation_mapping)
        return best_match if best_score > 0 else False

    def _score_depreciation_match(
        self,
        asset_code,
        asset_mapping,
        asset_summary,
        depreciation_code,
        depreciation_mapping,
        depreciation_summary,
    ):
        score = 0
        if self._asset_match_key(asset_code) == self._asset_match_key(depreciation_code):
            score += 80

        asset_moves = {sample["old_move_number"] for sample in asset_summary["samples"] if sample.get("old_move_number")}
        depreciation_moves = {sample["old_move_number"] for sample in depreciation_summary["samples"] if sample.get("old_move_number")}
        if asset_moves & depreciation_moves:
            score += 40

        asset_refs = self._extract_reference_tokens(asset_summary["samples"])
        depreciation_refs = self._extract_reference_tokens(depreciation_summary["samples"])
        if asset_refs & depreciation_refs:
            score += 30

        asset_text = self._normalize_keyword_text(
            " ".join(filter(None, [asset_mapping.get("source_account_name")] + [
                sample.get("description") or sample.get("concept") or "" for sample in asset_summary["samples"]
            ]))
        )
        depreciation_text = self._normalize_keyword_text(
            " ".join(filter(None, [depreciation_mapping.get("source_account_name")] + [
                sample.get("description") or sample.get("concept") or "" for sample in depreciation_summary["samples"]
            ]))
        )
        if asset_text and depreciation_text and self._significant_token_overlap(asset_text, depreciation_text):
            score += 20

        return score

    @staticmethod
    def _extract_reference_tokens(samples):
        tokens = set()
        for sample in samples:
            for source in filter(None, [sample.get("reference"), sample.get("description"), sample.get("concept")]):
                for match in re.findall(r"[A-Z]\d{1,3}-\d{1,4}|P\d{1,3}-\d{1,4}", source.upper()):
                    tokens.add(match)
        return tokens

    @staticmethod
    def _significant_token_overlap(asset_text, depreciation_text):
        asset_tokens = {token for token in asset_text.split() if len(token) > 3}
        depreciation_tokens = {token for token in depreciation_text.split() if len(token) > 3}
        return bool(asset_tokens & depreciation_tokens)

    @staticmethod
    def _pick_primary_sample(samples):
        return samples[0] if samples else {}

    def _get_expected_depreciation_target_code(self, old_code):
        depreciation_prefix = self.ASSET_DEPRECIATION_PREFIX_MAP.get(old_code[:3])
        if not depreciation_prefix:
            return False
        if len(old_code) > 4:
            return f"{depreciation_prefix}{old_code[4:]}"
        return self.DEPRECIATION_PREFIX_MAP.get(depreciation_prefix)

    def _asset_match_key(self, old_code):
        return old_code[-3:] if len(old_code) >= 3 else old_code

    def _guess_acquisition_date(self, samples):
        dates = [sample["date"] for sample in samples if sample.get("date")]
        return min(dates) if dates else False

    def _build_summary_text(self, diary_summary, account_count, partner_count, asset_count, tax_count, warning_count, error_count):
        return _(
            "Asientos detectados: %(moves)s\n"
            "Líneas detectadas: %(lines)s\n"
            "Cuentas antiguas detectadas: %(accounts)s\n"
            "Contactos detectados: %(partners)s\n"
            "Activos detectados: %(assets)s\n"
            "Impuestos detectados: %(taxes)s\n"
            "Advertencias: %(warnings)s\n"
            "Errores bloqueantes: %(errors)s"
        ) % {
            "moves": len(diary_summary["entries"]),
            "lines": sum(entry["count"] for entry in diary_summary["entries"].values()),
            "accounts": account_count,
            "partners": partner_count,
            "assets": asset_count,
            "taxes": tax_count,
            "warnings": warning_count,
            "errors": error_count,
        }

    def _build_review_report(self, summary_text, mapping_lines, asset_lines, move_preview_lines, log_lines):
        parts = [summary_text, "", "MAPEOS"]
        for line in mapping_lines:
            parts.append(
                "- %(source)s -> %(target)s [%(mapping)s] %(partner)s %(warning)s" % {
                    "source": line["source_code"],
                    "target": line["target_account_code"],
                    "mapping": line["mapping_type"],
                    "partner": line.get("partner_name") or "",
                    "warning": line.get("warning_message") or "",
                }
            )
        parts.append("")
        parts.append("ACTIVOS")
        for line in asset_lines:
            parts.append(
                "- %(source)s | %(name)s | %(target)s | valor %(value)s | amortización %(depr)s -> %(depr_target)s | acumulada %(accumulated)s | pendiente %(residual)s | %(state)s | %(warning)s" % {
                    "source": line["source_code"],
                    "name": line.get("asset_name") or "",
                    "target": line["target_account_id"] and line["target_account_id"] or "",
                    "value": line.get("original_value") or 0,
                    "depr": line.get("depreciation_source_code") or "-",
                    "depr_target": line.get("depreciation_target_account_id") or "-",
                    "accumulated": line.get("accumulated_value") or 0,
                    "residual": line.get("residual_value") or 0,
                    "state": line.get("state") or "review_required",
                    "warning": line.get("warning_message") or "",
                }
            )
        parts.append("")
        parts.append("ASIENTOS")
        for line in move_preview_lines:
            parts.append(
                "- %(move)s %(date)s %(status)s Debe:%(debit)s Haber:%(credit)s" % {
                    "move": line["old_move_number"],
                    "date": fields.Date.to_string(line["move_date"]),
                    "status": line["status"],
                    "debit": line["debit_total"],
                    "credit": line["credit_total"],
                }
            )
        parts.append("")
        parts.append("LOGS")
        for log in log_lines:
            parts.append("- [%(level)s/%(category)s] %(message)s" % log)
        return "\n".join(parts)

    def _make_log_line(self, level, category, message, old_account_code=False, move_number=False, source_file=False, source_line=False, blocking=False):
        return {
            "level": level,
            "category": category,
            "message": message,
            "old_account_code": old_account_code,
            "move_number": move_number,
            "source_file": source_file,
            "source_line": source_line,
            "blocking": blocking,
        }

    def _reset_review_lines(self):
        self.write({
            "mapping_line_ids": [Command.clear()],
            "asset_line_ids": [Command.clear()],
            "move_preview_line_ids": [Command.clear()],
            "log_line_ids": [Command.clear()],
        })

    def _load_excel_rows(self, file_content, file_name, sheet_name=False):
        file_data = base64.b64decode(file_content)
        is_xls = file_name and file_name.lower().endswith(".xls")
        try:
            if is_xls:
                return self._read_xls(file_data, sheet_name)
            return self._read_xlsx(file_data, sheet_name)
        except Exception as exc:
            raise UserError(_("Fichero Excel no válido: %s") % str(exc)) from exc

    def _read_xlsx(self, file_data, sheet_name=False):
        workbook = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        return list(sheet.iter_rows(values_only=True))

    def _read_xls(self, file_data, sheet_name=False):
        workbook = xlrd.open_workbook(file_contents=file_data)
        sheet_names = workbook.sheet_names()
        sheet = workbook.sheet_by_name(sheet_name) if sheet_name and sheet_name in sheet_names else workbook.sheet_by_index(0)
        rows = []
        for row_idx in range(sheet.nrows):
            row_vals = []
            for col_idx in range(sheet.ncols):
                cell = sheet.cell(row_idx, col_idx)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    dt_tuple = xlrd.xldate_as_tuple(cell.value, workbook.datemode)
                    row_vals.append(datetime(*dt_tuple))
                else:
                    row_vals.append(cell.value)
            rows.append(row_vals)
        return rows

    def _locate_diary_columns(self, rows):
        for row_index, row in enumerate(rows):
            columns = {}
            for cell_index, value in enumerate(row):
                key = self.DIARY_HEADER_ALIASES.get(self._normalize_header(value))
                if key and key not in columns:
                    columns[key] = cell_index
            if self.DIARY_REQUIRED_COLUMNS.issubset(columns):
                return row_index, columns
        expected = ", ".join(sorted(self.DIARY_REQUIRED_COLUMNS | self.DIARY_OPTIONAL_COLUMNS))
        raise UserError(
            _("No se ha encontrado la fila de cabecera del Libro Diario. Columnas esperadas: %s")
            % expected
        )

    def _find_saved_mapping(self, source_code):
        return self.env["kolac.account.subaccount.map"].search([
            ("company_id", "=", self.company_id.id),
            ("source_code", "=", source_code),
        ], limit=1)

    def _get_account_model(self):
        return self.env["account.account"].with_company(self.company_id).with_context(
            allowed_company_ids=self.company_id.ids,
            active_test=False,
        )

    def _find_account(self, code):
        if not code:
            return self.env["account.account"]
        return self._get_account_model().search([("code", "=", code)], limit=1)

    def _create_missing_account(self, code, description="", mapping_type=False):
        existing = self._find_account(code)
        if existing:
            return existing
        create_vals = self._get_missing_account_vals(code, description, mapping_type=mapping_type)
        if not create_vals:
            return self.env["account.account"]
        account = self._get_account_model().sudo().create(create_vals)
        return self._get_account_model().browse(account.id)

    def _get_missing_account_vals(self, code, description="", mapping_type=False):
        if not code or len(code) < 3:
            return {}
        template = self._find_template_account(code)
        group = self._find_matching_group(code)
        default_vals = self._get_missing_account_defaults(code)
        if not template and not group and not default_vals and not mapping_type:
            return {}
        vals = {
            "code": code,
            "name": (
                default_vals.get("name")
                or description
                or (template and template.name)
                or (group and group.name)
                or _("Cuenta %s") % code
            )[:255],
            "company_ids": [Command.set(self.company_id.ids)],
        }
        if template:
            vals.update({
                "account_type": template.account_type,
                "reconcile": template.reconcile,
                "tag_ids": [Command.set(template.tag_ids.ids)],
            })
            if "currency_id" in template._fields and template.currency_id:
                vals["currency_id"] = template.currency_id.id
            if "non_trade" in template._fields:
                vals["non_trade"] = template.non_trade
        elif default_vals:
            vals.update(default_vals)
        elif mapping_type:
            vals.update(self._get_mapping_type_account_vals(mapping_type))
        if "account_type" not in vals:
            inferred_type = self._infer_account_type_from_code(code)
            if inferred_type:
                vals["account_type"] = inferred_type
        return vals

    @staticmethod
    def _infer_account_type_from_code(code):
        if not code:
            return False
        first = str(code)[0]
        type_by_prefix = {
            "1": "equity",
            "2": "asset_fixed",
            "3": "asset_current",
            "4": "liability_payable",
            "5": "asset_cash",
            "6": "expense",
            "7": "income",
            "8": "equity",
            "9": "equity",
        }
        return type_by_prefix.get(first)

    def _can_auto_create_missing_account(self, code):
        if not code or len(code) < 3:
            return False
        if self.create_missing_accounts:
            return True
        return bool(
            self._get_missing_account_defaults(code)
            or self._find_template_account(code)
            or self._find_matching_group(code)
        )

    def _get_missing_account_defaults(self, code):
        for candidate in self._get_template_code_candidates(code):
            defaults = self.MISSING_ACCOUNT_DEFAULTS.get(candidate)
            if defaults:
                return defaults
        return {}

    @staticmethod
    def _get_mapping_type_account_vals(mapping_type):
        mapping_defaults = {
            "partner_customer": {"account_type": "asset_receivable", "reconcile": True},
            "partner_supplier": {"account_type": "liability_payable", "reconcile": True},
            "partner_creditor": {"account_type": "liability_payable", "reconcile": True},
            "grant_or_becario_payable": {"account_type": "liability_payable", "reconcile": True},
            "asset": {"account_type": "asset_fixed"},
            "accumulated_depreciation": {"account_type": "asset_non_current"},
            "tax_input": {"account_type": "asset_current"},
            "tax_output": {"account_type": "liability_current"},
            "bank": {"account_type": "asset_cash"},
            "cash": {"account_type": "asset_cash"},
            "expense_nature": {"account_type": "expense"},
            "expense_nature_by_supplier": {"account_type": "expense"},
            "direct_account": {},
            "review_required": {},
        }
        return dict(mapping_defaults.get(mapping_type, {}))

    def _find_template_account(self, code):
        candidate = self._find_account(code)
        if candidate:
            return candidate
        for candidate_code in self._get_template_code_candidates(code):
            candidate = self._find_account(candidate_code)
            if candidate:
                return candidate
        prefix_template = self._find_prefix_template_account(code)
        if prefix_template:
            return prefix_template
        group = self._find_matching_group(code)
        if group:
            return self._find_group_anchor_account(code, group)
        return self.env["account.account"]

    def _get_template_code_candidates(self, code):
        candidates = []
        if code and len(code) >= 6:
            candidates.append(code[:6])
        if code[:3] in self.PARTNER_PREFIX_MAP:
            candidates.append(f"{code[:3]}000")
        if code[:3] in self.ASSET_PREFIX_MAP:
            candidates.append(self.ASSET_PREFIX_MAP[code[:3]])
        if code[:4] in self.DEPRECIATION_PREFIX_MAP:
            candidates.append(self.DEPRECIATION_PREFIX_MAP[code[:4]])
        if code[:3] in self.TAX_PREFIX_MAP:
            candidates.append(f"{code[:3]}000")
        if code.startswith("572"):
            candidates.append("572000")
        if code.startswith("570"):
            candidates.append("570000")
        seen = set()
        return [candidate for candidate in candidates if candidate and not (candidate in seen or seen.add(candidate))]

    def _find_prefix_template_account(self, code):
        account_model = self._get_account_model()
        candidates = self.env["account.account"]
        for prefix_length in range(len(code) - 1, 2, -1):
            candidate = account_model.search([("code", "=", code[:prefix_length])], limit=1)
            if candidate:
                candidates |= candidate
        if candidates:
            return candidates.sorted(
                lambda account: (
                    len(account.code or ""),
                    -(len((account.code or "")) - len((account.code or "").rstrip("0"))),
                    account.code or "",
                )
            )[:1]
        return self.env["account.account"]

    def _get_group_model(self):
        return self.env["account.group"].with_context(active_test=False)

    def _find_matching_group(self, code):
        groups = self._get_group_model().search([("company_id", "=", self.company_id.root_id.id)])
        matching_groups = groups.filtered(lambda group: self._group_matches_code(group, code))
        if not matching_groups:
            return self.env["account.group"]
        return matching_groups.sorted(lambda group: len(group.code_prefix_start or ""), reverse=True)[:1]

    def _find_group_anchor_account(self, code, group):
        all_accounts = self._get_account_model().search([])
        current_group = group
        while current_group:
            candidates = all_accounts.filtered(
                lambda account: self._group_matches_code(current_group, account.code)
            )
            if candidates:
                return candidates.sorted(lambda account: account.code or "")[:1]
            current_group = current_group.parent_id
        return self.env["account.account"]

    @staticmethod
    def _group_matches_code(group, code):
        if not group.code_prefix_start or not code:
            return False
        prefix_length = len(group.code_prefix_start)
        code_prefix = code[:prefix_length]
        code_prefix_end = group.code_prefix_end or group.code_prefix_start
        return group.code_prefix_start <= code_prefix <= code_prefix_end

    @classmethod
    def _is_zero_amount_row(cls, debit, credit):
        return cls._round_amount(debit) == Decimal("0.00") and cls._round_amount(credit) == Decimal("0.00")

    @classmethod
    def _is_summary_row(cls, row, entry_date, old_account_code):
        row_text = cls._normalize_keyword_text(" ".join(
            cls._cell_to_text(value) for value in row if value not in (None, False, "")
        ))
        if not row_text:
            return False
        strong_keywords = (
            "descuadre",
            "diferencia",
            "total contado",
            "total contada",
            "cantidad total contado",
            "cantidad total contada",
            "importe total",
            "total diario",
        )
        weak_keywords = ("total", "subtotal", "saldo", "resumen")
        if any(keyword in row_text for keyword in strong_keywords):
            return not entry_date or not old_account_code
        if any(keyword in row_text for keyword in weak_keywords):
            return not entry_date and not old_account_code
        return False

    @staticmethod
    def _normalize_keyword_text(value):
        text = KolacAccountJournalImportWizard._cell_to_text(value).lower()
        if not text:
            return ""
        text = "".join(
            char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char)
        )
        return re.sub(r"\s+", " ", text).strip()

    @staticmethod
    def _sanitize_old_account_code(value):
        return re.sub(r"[^0-9A-Za-z]", "", KolacAccountJournalImportWizard._cell_to_text(value))

    @staticmethod
    def _get_value(row, columns, key):
        index = columns.get(key)
        if index is None or index >= len(row):
            return None
        return row[index]

    @staticmethod
    def _normalize_header(value):
        text = KolacAccountJournalImportWizard._cell_to_text(value).lower()
        if not text:
            return ""
        text = "".join(
            char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char)
        )
        return re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    @staticmethod
    def _cell_to_text(value):
        if value in (None, False):
            return ""
        if isinstance(value, str):
            return value.strip()
        if isinstance(value, float):
            text = format(value, "f").rstrip("0").rstrip(".")
            return text or ""
        if isinstance(value, int):
            return str(value)
        if isinstance(value, (datetime, date)):
            return fields.Date.to_string(value.date() if isinstance(value, datetime) else value)
        return str(value).strip()

    def _parse_date(self, value, default_year=False):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = self._cell_to_text(value)
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        numeric_match = re.match(
            r"^(?P<day>\d{1,2})[-/](?P<month>\d{1,2})(?:[-/](?P<year>\d{2,4}))?$",
            text,
        )
        if numeric_match:
            month = int(numeric_match.group("month"))
            year = numeric_match.group("year")
            if 1 <= month <= 12 and (year or default_year):
                resolved_year = int(year) if year else int(default_year)
                if resolved_year < 100:
                    resolved_year += 2000
                try:
                    return date(resolved_year, month, int(numeric_match.group("day")))
                except ValueError:
                    return False
        month_match = re.match(
            r"^(?P<day>\d{1,2})[-/ ](?P<month>[A-Za-zÁÉÍÓÚÜÑáéíóúüñ.]+)(?:[-/ ](?P<year>\d{2,4}))?\.?$",
            text,
        )
        if month_match:
            month_key = self._normalize_keyword_text(month_match.group("month")).replace(".", "")[:3]
            month_map = {
                "ene": 1,
                "jan": 1,
                "feb": 2,
                "mar": 3,
                "abr": 4,
                "apr": 4,
                "may": 5,
                "jun": 6,
                "jul": 7,
                "ago": 8,
                "aug": 8,
                "sep": 9,
                "oct": 10,
                "nov": 11,
                "dic": 12,
                "dec": 12,
            }
            month = month_map.get(month_key)
            year = month_match.group("year")
            if month and (year or default_year):
                resolved_year = int(year) if year else int(default_year)
                if resolved_year < 100:
                    resolved_year += 2000
                return date(resolved_year, month, int(month_match.group("day")))
        try:
            return fields.Date.to_date(text)
        except Exception:
            return False

    @staticmethod
    def _parse_amount(value):
        if value in (None, False, ""):
            return Decimal("0.00")
        if isinstance(value, Decimal):
            return KolacAccountJournalImportWizard._round_amount(value)
        if isinstance(value, (int, float)):
            return KolacAccountJournalImportWizard._round_amount(Decimal(str(value)))
        text = KolacAccountJournalImportWizard._cell_to_text(value)
        if not text:
            return Decimal("0.00")
        text = text.replace(" ", "").replace("\xa0", "")
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1]
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        try:
            amount = Decimal(text)
        except InvalidOperation as exc:
            raise UserError(_("No se ha podido interpretar el importe '%s'.") % value) from exc
        if negative:
            amount *= Decimal("-1")
        return KolacAccountJournalImportWizard._round_amount(amount)

    @staticmethod
    def _round_amount(amount):
        return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @staticmethod
    def _is_blank_row(row):
        return not any(value not in (None, False, "") and str(value).strip() for value in row)

    def _reopen_wizard(self):
        return {
            "type": "ir.actions.act_window",
            "name": _("Importación contable KOLAC"),
            "res_model": self._name,
            "view_mode": "form",
            "target": "new",
            "res_id": self.id,
        }

    def _open_created_moves(self, moves):
        action = self.env["ir.actions.act_window"]._for_xml_id("account.action_move_journal_line")
        action["domain"] = [("id", "in", moves.ids)]
        action["context"] = {
            "default_move_type": "entry",
            "default_journal_id": self.journal_id.id,
            "search_default_posted": 0,
            "search_default_unposted": 0,
            "view_no_maturity": True,
        }
        if len(moves) == 1:
            action["view_mode"] = "form"
            action["res_id"] = moves.id
        return action

    def _is_move_balanced(self, move):
        total_debit = sum(move.line_ids.mapped("debit"))
        total_credit = sum(move.line_ids.mapped("credit"))
        return self._round_amount(Decimal(str(total_debit))) == self._round_amount(Decimal(str(total_credit)))


class KolacAccountJournalImportMapLine(models.TransientModel):
    _name = "kolac.account.journal.import.map.line"
    _description = "Kolac Import Mapping Review"

    def init(self):
        self.env.cr.execute(
            """
            UPDATE kolac_account_journal_import_map_line
               SET mapping_type = 'review_required'
             WHERE mapping_type IS NULL
            """
        )

    wizard_id = fields.Many2one("kolac.account.journal.import.wizard", required=True, ondelete="cascade")
    source_code = fields.Char(string="Cuenta antigua", required=True, readonly=True)
    source_account_name = fields.Char(string="Nombre detectado")
    source_origin = fields.Selection(
        [("journal", "Libro diario"), ("account_list", "Listado"), ("account_plan", "Plan"), ("manual", "Manual")],
        string="Origen",
        readonly=True,
    )
    account_list_name = fields.Char(string="Nombre listado", readonly=True)
    plan_account_name = fields.Char(string="Nombre plan", readonly=True)
    occurrence_count = fields.Integer(string="Ocurrencias", readonly=True)
    saved_mapping_id = fields.Many2one("kolac.account.subaccount.map", string="Mapeo guardado", readonly=True)
    mapping_type = fields.Selection(KolacAccountJournalImportWizard.MAPPING_TYPES, string="Tipo de mapeo", required=True)
    target_account_code = fields.Char(string="Código final", readonly=True)
    company_id = fields.Many2one(related="wizard_id.company_id", readonly=True)
    target_account_id = fields.Many2one(
        "account.account",
        string="Cuenta Odoo",
        domain="[('company_ids', 'in', [company_id])]",
    )
    partner_name = fields.Char(string="Contacto detectado")
    partner_id = fields.Many2one("res.partner", string="Contacto")
    target_tax_id = fields.Many2one("account.tax", string="Impuesto sugerido")
    if _ACCOUNT_ASSET_AVAILABLE:
        asset_id = fields.Many2one("account.asset", string="Activo existente")
    suggestion_origin = fields.Selection(
        [("saved", "Guardado"), ("suggested", "Inferido"), ("none", "Sin sugerencia")],
        string="Origen sugerencia",
        readonly=True,
    )
    action = fields.Selection(
        [("reuse", "Reutilizar"), ("create_partner", "Crear contacto"), ("create_asset", "Crear activo"), ("review", "Revisar")],
        string="Acción",
        default="reuse",
    )
    save_mapping = fields.Boolean(string="Guardar mapeo", default=True)
    review_required = fields.Boolean(string="Revisión requerida")
    warning_message = fields.Text(string="Advertencias")
    debit_total = fields.Monetary(string="Debe total", currency_field="company_currency_id", readonly=True)
    credit_total = fields.Monetary(string="Haber total", currency_field="company_currency_id", readonly=True)
    company_currency_id = fields.Many2one(related="wizard_id.company_id.currency_id", readonly=True)


class KolacAccountJournalImportAssetLine(models.TransientModel):
    _name = "kolac.account.journal.import.asset.line"
    _description = "Kolac Asset Review"

    wizard_id = fields.Many2one("kolac.account.journal.import.wizard", required=True, ondelete="cascade")
    source_code = fields.Char(string="Cuenta activo antigua", readonly=True)
    asset_name = fields.Char(string="Activo detectado")
    old_account_name = fields.Char(string="Nombre antiguo", readonly=True)
    target_account_id = fields.Many2one("account.account", string="Cuenta activo Odoo", readonly=True)
    depreciation_source_code = fields.Char(string="Cuenta amortización antigua", readonly=True)
    old_depreciation_account_name = fields.Char(string="Nombre amortización antigua", readonly=True)
    depreciation_target_account_id = fields.Many2one("account.account", string="Cuenta amortización Odoo", readonly=True)
    original_value = fields.Monetary(string="Valor original", currency_field="company_currency_id", readonly=True)
    accumulated_value = fields.Monetary(string="Amortización acumulada", currency_field="company_currency_id", readonly=True)
    residual_value = fields.Monetary(string="Pendiente", currency_field="company_currency_id", readonly=True)
    acquisition_date = fields.Date(string="Fecha adquisición")
    create_asset = fields.Boolean(string="Crear activo")
    if _ACCOUNT_ASSET_AVAILABLE:
        asset_id = fields.Many2one("account.asset", string="Activo Odoo")
    state = fields.Selection(
        [
            ("mapped_with_asset", "Listo para crear activo"),
            ("fully_depreciated", "Totalmente amortizado"),
            ("review_required", "Revisión requerida"),
            ("error", "Error contable"),
        ],
        string="Estado",
        readonly=True,
    )
    source_move_number = fields.Char(string="Asiento origen", readonly=True)
    source_reference = fields.Char(string="Referencia origen", readonly=True)
    source_label = fields.Char(string="Descripción original", readonly=True)
    source_file = fields.Char(string="Fichero origen", readonly=True)
    warning_message = fields.Text(string="Advertencias")
    company_currency_id = fields.Many2one(related="wizard_id.company_id.currency_id", readonly=True)


class KolacAccountJournalImportMovePreview(models.TransientModel):
    _name = "kolac.account.journal.import.move.preview"
    _description = "Kolac Move Preview"

    wizard_id = fields.Many2one("kolac.account.journal.import.wizard", required=True, ondelete="cascade")
    old_move_number = fields.Char(string="Asiento origen", readonly=True)
    move_date = fields.Date(string="Fecha", readonly=True)
    reference = fields.Char(string="Referencia", readonly=True)
    concept = fields.Char(string="Concepto", readonly=True)
    line_count = fields.Integer(string="Líneas", readonly=True)
    debit_total = fields.Monetary(string="Debe", currency_field="company_currency_id", readonly=True)
    credit_total = fields.Monetary(string="Haber", currency_field="company_currency_id", readonly=True)
    target_journal_id = fields.Many2one("account.journal", string="Diario destino", domain="[('company_id', '=', wizard_id.company_id)]")
    is_balanced = fields.Boolean(string="Cuadrado", readonly=True)
    status = fields.Selection([("ready", "Preparado"), ("error", "Error")], string="Estado", readonly=True)
    company_currency_id = fields.Many2one(related="wizard_id.company_id.currency_id", readonly=True)


class KolacAccountJournalImportLogLine(models.TransientModel):
    _name = "kolac.account.journal.import.log.line"
    _description = "Kolac Import Log"

    wizard_id = fields.Many2one("kolac.account.journal.import.wizard", required=True, ondelete="cascade")
    level = fields.Selection([("info", "Info"), ("warning", "Advertencia"), ("error", "Error")], string="Nivel", readonly=True)
    category = fields.Char(string="Categoría", readonly=True)
    message = fields.Text(string="Mensaje", readonly=True)
    old_account_code = fields.Char(string="Cuenta antigua", readonly=True)
    move_number = fields.Char(string="Asiento origen", readonly=True)
    source_file = fields.Char(string="Fichero", readonly=True)
    source_line = fields.Integer(string="Línea", readonly=True)
    blocking = fields.Boolean(string="Bloqueante", readonly=True)
